from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from app import timeutil
from app.enums import FREQUENCY_LABELS, REQUEST_TYPE_LABELS, STATUS_LABELS
from app.models import Request, User
from app.services.calendar import BusinessCalendar
from app.services.requests import SORT_COLUMNS, ListFilters, filtered_query, topic_display

MAX_ROWS = 20000

COLUMNS = [
    ("Номер", 18), ("Подана", 17), ("Статус", 22), ("Тип", 20), ("Тематика", 20), ("Заявитель", 30), ("Подразделение", 30),
    ("Регион", 18), ("Телефон", 22), ("Почта", 26), ("Рассмотреть до", 15), ("Срок реализации", 15), ("Ответственный", 28),
    ("Периодичность", 14), ("Трудоёмкость, ч/мес", 14), ("Текущий процесс", 60), ("Что автоматизировать", 60), ("Желаемый результат", 60),
]


def _text(value: str | None) -> str | None:
    # Строка, начинающаяся с «=», в Excel станет формулой. Экранируем, чтобы текст заявки не исполнялся.
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


def export_xlsx(db: Session, user: User, filters: ListFilters) -> bytes:
    calendar = BusinessCalendar.load(db)
    today = timeutil.local_today()
    stmt = filtered_query(db, user, filters, calendar, today)
    column = SORT_COLUMNS.get(filters.sort, Request.submitted_at)
    stmt = stmt.options(selectinload(Request.topic), selectinload(Request.region)).order_by(
        column.asc() if filters.direction == "asc" else column.desc(), Request.submitted_at.desc()
    ).limit(MAX_ROWS)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заявки"
    sheet.append([title for title, _ in COLUMNS])
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.cell(row=1, column=index).font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for request in db.scalars(stmt):
        sheet.append(
            [
                request.number,
                timeutil.to_local(request.submitted_at).replace(tzinfo=None),
                STATUS_LABELS[request.status],
                REQUEST_TYPE_LABELS[request.type],
                _text(topic_display(request)),
                _text(request.applicant_name),
                _text(request.applicant_department),
                request.region.name,
                _text(request.applicant_phone),
                _text(request.applicant_email),
                request.review_due_date,
                request.deadline,
                _text(request.responsible),
                FREQUENCY_LABELS[request.frequency],
                float(request.workload_hours_month) if request.workload_hours_month is not None else None,
                _text(request.process),
                _text(request.problem),
                _text(request.desired_result),
            ]
        )
        row = sheet.max_row
        sheet.cell(row=row, column=2).number_format = "DD.MM.YYYY HH:MM"
        sheet.cell(row=row, column=11).number_format = "DD.MM.YYYY"
        sheet.cell(row=row, column=12).number_format = "DD.MM.YYYY"
        for col in (16, 17, 18):
            sheet.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
