"""Сценарии через HTTP API. Номера в комментариях — пункты аудита прототипа."""

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.services.documents import storage_path
from tests.conftest import TODAY, TOPIC_OTHER, create_request, errors_by_path, payload


def transition(api, request, to_status, **extra):
    return api.post(f"/api/requests/{request['id']}/transitions", json={"to_status": to_status, "version": request["version"], **extra})


def accept(api, request, **extra):
    body = {"responsible": "Группа развития ИС", "deadline": (TODAY + timedelta(days=30)).isoformat(), **extra}
    response = transition(api, request, "accepted", **body)
    assert response.status_code == 200, response.text
    return response.json()


# ---------------------------------------------------------------- подача

def test_create_request_assigns_number_deadline_and_notifies_admins(applicant, admin):
    request = create_request(applicant)

    assert request["number"] == "AUT-2026-000001"
    assert request["status"] == "review"
    assert request["applicant_name"] == "Смирнова Ольга Викторовна"
    assert request["applicant_phone"] == "+7 (912) 345-67-89"
    assert request["review_due_date"] == "2026-09-21"
    assert request["sla"] == {"state": "ok", "review_due_date": "2026-09-21", "business_days_left": 5}
    assert float(request["workload_hours_month"]) == 3.0
    assert [event["kind"] for event in request["events"]] == ["submitted"]
    assert [action["kind"] for action in request["actions"]] == ["withdraw"]

    notifications = admin.get("/api/notifications").json()
    assert notifications["unread"] == 1
    assert "AUT-2026-000001" in notifications["items"][0]["text"]
    assert applicant.get("/api/me").json()["unread_notifications"] == 0


def test_applicant_defaults_are_remembered(applicant):
    create_request(applicant, applicant={"phone": "+7 900 111-22-33", "region_id": 2})
    defaults = applicant.get("/api/me").json()["applicant_defaults"]
    assert defaults["phone"] == "+7 (900) 111-22-33"
    assert defaults["region_id"] == 2
    assert defaults["full_name"] == "Смирнова Ольга Викторовна"


def test_field_errors_are_russian_and_point_to_fields(applicant):
    # Л-1, Л-13: отрицательные значения и мусорный телефон отклоняются сервером.
    response = applicant.post(
        "/api/requests",
        json=payload(applicant={"phone": "0000000000"}, workload={"duration": "-100", "times_per_period": -5, "employees_count": 0}),
    )
    assert response.status_code == 422
    errors = errors_by_path(response)
    assert errors["workload.duration"] == "Значение должно быть больше 0"
    assert errors["workload.times_per_period"] == "Значение должно быть не меньше 1"
    assert errors["workload.employees_count"] == "Значение должно быть не меньше 1"
    assert "+7 912 345-67-89" in errors["applicant.phone"]


def test_cross_field_rules_are_reported_together(applicant):
    response = applicant.post(
        "/api/requests",
        json=payload(
            type="upgrade",
            topic_id=TOPIC_OTHER,
            workload={"frequency": "week", "duration": None, "duration_unit": None, "times_per_period": None, "employees_count": None},
        ),
    )
    assert response.status_code == 422
    assert set(errors_by_path(response)) == {
        "existing_system_name",
        "topic_other",
        "workload.duration",
        "workload.duration_unit",
        "workload.times_per_period",
        "workload.employees_count",
    }


def test_consent_is_required(applicant):
    response = applicant.post("/api/requests", json=payload(consent=False))
    assert response.status_code == 422
    assert set(errors_by_path(response)) == {"consent"}


def test_irregular_process_needs_description_but_no_numbers(applicant):
    response = applicant.post("/api/requests", json=payload(workload={"frequency": "irregular", "duration": None, "times_per_period": None}))
    assert errors_by_path(response) == {"workload.note": "Опишите, как часто возникает процесс и сколько времени он занимает"}

    request = create_request(
        applicant, workload={"frequency": "irregular", "duration": None, "times_per_period": None, "note": "При каждой внеплановой проверке"}
    )
    assert request["workload_hours_month"] is None


def test_blank_optional_text_is_stored_as_empty(applicant):
    request = create_request(applicant, method_suggestion="   ", existing_system_name="")
    assert request["method_suggestion"] is None
    assert request["existing_system_name"] is None


# ---------------------------------------------------------------- доступ, поиск, сортировка

def test_applicant_sees_only_own_requests(applicant, other_applicant, admin):
    # Д-1: раньше «Мои заявки» показывали заявки всех сотрудников.
    mine = create_request(applicant)
    create_request(applicant)
    create_request(other_applicant)

    assert applicant.get("/api/requests").json()["total"] == 2
    assert other_applicant.get("/api/requests").json()["total"] == 1
    assert other_applicant.get("/api/requests?scope=all").json()["total"] == 1
    assert other_applicant.get(f"/api/requests/{mine['id']}").status_code == 404
    assert admin.get("/api/requests?scope=all").json()["total"] == 3


def test_search_finds_words_in_any_order_and_case(applicant, other_applicant, admin):
    # БЛ-3: поиск в прототипе не работал совсем и искал только по номеру.
    create_request(applicant)
    create_request(other_applicant, process="Диспетчер оформляет путевые листы водителей на бумаге.")

    def found(query):
        return [item["number"] for item in admin.get("/api/requests", params={"scope": "all", "q": query}).json()["items"]]

    assert found("ЛИСТЫ путевые") == ["AUT-2026-000002"]
    assert found("ковалёв") == ["AUT-2026-000002"]
    assert found("AUT-2026-000001") == ["AUT-2026-000001"]
    assert found("%") == []
    assert found("нет-такого-слова") == []


def test_sorting_and_pagination(applicant, admin):
    # БЛ-4 и UX-8: сортировка по любой колонке и постраничный вывод.
    for _ in range(3):
        create_request(applicant)

    def numbers(**params):
        return [item["number"] for item in admin.get("/api/requests", params={"scope": "all", **params}).json()["items"]]

    assert numbers(sort="number", direction="asc") == ["AUT-2026-000001", "AUT-2026-000002", "AUT-2026-000003"]
    assert numbers(sort="number", direction="desc") == ["AUT-2026-000003", "AUT-2026-000002", "AUT-2026-000001"]

    page = admin.get("/api/requests", params={"scope": "all", "sort": "number", "direction": "asc", "page": 2, "page_size": 2}).json()
    assert page["total"] == 3
    assert [item["number"] for item in page["items"]] == ["AUT-2026-000003"]

    for field in ["status", "applicant", "department", "region", "topic", "type", "review_due_date", "deadline", "responsible", "updated_at"]:
        assert admin.get("/api/requests", params={"scope": "all", "sort": field}).status_code == 200
    assert admin.get("/api/requests", params={"sort": "drop table"}).status_code == 422


def test_filters_by_status_topic_region(applicant, admin):
    first = create_request(applicant)
    create_request(applicant, topic_id=TOPIC_OTHER, topic_other="Социальные программы", applicant={"region_id": 3})
    transition(admin, first, "rejected", comment="Нецелесообразно")

    def total(**params):
        return admin.get("/api/requests", params={"scope": "all", **params}).json()["total"]

    assert total(status="rejected") == 1
    assert total(status=["rejected", "review"]) == 2
    assert total(topic_id=TOPIC_OTHER) == 1
    assert total(region_id=3) == 1


# ---------------------------------------------------------------- действия администратора

def test_accept_requires_responsible_and_future_deadline(applicant, admin):
    # БЛ-1, Л-15, Л-16: принятие в работу теперь выполнимо и проверяет срок.
    request = create_request(applicant)

    response = transition(admin, request, "accepted")
    assert set(errors_by_path(response)) == {"responsible", "deadline"}

    response = transition(admin, request, "accepted", responsible="ИТ", deadline=(TODAY - timedelta(days=1)).isoformat())
    assert errors_by_path(response) == {"deadline": "Срок не может быть в прошлом"}

    accepted = accept(admin, request, comment="Берём в работу")
    assert accepted["status"] == "accepted"
    assert accepted["responsible"] == "Группа развития ИС"
    assert accepted["review_due_date"] is None
    assert accepted["decided_at"] is not None
    assert accepted["events"][-1]["actor_name"] == "Орлов Сергей Николаевич"
    assert applicant.get("/api/notifications").json()["unread"] == 1


def test_applicant_cannot_change_status(applicant):
    request = create_request(applicant)
    response = transition(applicant, request, "accepted", responsible="Я", deadline="2030-01-01")
    assert response.status_code == 403


def test_stale_version_is_rejected(applicant, admin, as_user):
    request = create_request(applicant)
    accept(admin, request)
    response = transition(as_user("s.orlov"), request, "rejected", comment="Параллельное решение")
    assert response.status_code == 409
    assert response.json()["code"] == "stale_version"


def test_impossible_transition_is_rejected(applicant, admin):
    request = create_request(applicant)
    response = transition(admin, request, "done")
    assert response.status_code == 409
    assert response.json()["code"] == "invalid_transition"


def test_full_lifecycle_with_rollback_and_close(applicant, admin):
    request = accept(admin, create_request(applicant))
    request = transition(admin, request, "development").json()

    # Л-17: ошибочный перевод можно откатить, но только с объяснением.
    assert errors_by_path(transition(admin, request, "accepted")) == {"comment": "Укажите причину — заявитель увидит этот комментарий"}
    request = transition(admin, request, "accepted", comment="Нужно согласовать ТЗ").json()
    assert request["status"] == "accepted"

    for status in ["development", "testing", "done", "closed"]:
        response = transition(admin, request, status)
        assert response.status_code == 200, response.text
        request = response.json()
    assert request["closed_at"] is not None
    assert request["actions"] == []


def test_rejected_request_can_be_reopened(applicant, admin):
    request = create_request(applicant)
    assert errors_by_path(transition(admin, request, "rejected")) == {"comment": "Укажите причину — заявитель увидит этот комментарий"}
    request = transition(admin, request, "rejected", comment="Мала трудоёмкость").json()

    reopened = transition(admin, request, "review", comment="Появились новые данные о трудоёмкости")
    assert reopened.status_code == 200
    assert reopened.json()["status"] == "review"
    assert reopened.json()["review_due_date"] == "2026-09-21"


def test_assignment_can_be_changed_while_in_work(applicant, admin):
    # Л-18: срок и ответственного можно менять после принятия.
    request = accept(admin, create_request(applicant))
    new_deadline = (TODAY + timedelta(days=60)).isoformat()

    response = admin.put(f"/api/requests/{request['id']}/assignment", json={"responsible": "Группа развития ИС", "deadline": new_deadline, "version": request["version"]})
    assert response.status_code == 200, response.text
    updated = response.json()
    assert updated["deadline"] == new_deadline
    assert updated["events"][-1]["kind"] == "assignment_changed"
    assert updated["events"][-1]["payload"]["deadline"][1] == new_deadline

    same = admin.put(f"/api/requests/{request['id']}/assignment", json={"responsible": "Группа развития ИС", "deadline": new_deadline, "version": updated["version"]})
    assert same.status_code == 422
    forbidden = applicant.put(f"/api/requests/{request['id']}/assignment", json={"responsible": "Я", "deadline": new_deadline, "version": updated["version"]})
    assert forbidden.status_code == 403


def test_admin_opening_request_marks_it_seen(applicant, admin):
    # Л-8: «новые» — это не просмотренные, а не «с двумя записями в истории».
    request = create_request(applicant)
    assert admin.get("/api/requests", params={"scope": "all", "attention": "unviewed"}).json()["total"] == 1

    detail = admin.get(f"/api/requests/{request['id']}").json()
    assert detail["first_viewed_at"] is not None
    assert detail["version"] == request["version"]
    assert admin.get("/api/requests", params={"scope": "all", "attention": "unviewed"}).json()["total"] == 0
    applicant.get(f"/api/requests/{request['id']}")
    assert admin.get("/api/admin/dashboard").json()["attention"]["unviewed"] == 0


# ---------------------------------------------------------------- уточнение, отзыв, комментарии

def test_clarification_pauses_deadline_and_resubmission_restarts_it(applicant, admin, freeze):
    # Л-7, Д-5, UX-14.
    request = create_request(applicant)
    assert errors_by_path(transition(admin, request, "clarification")) == {"comment": "Опишите, что нужно уточнить"}
    request = transition(admin, request, "clarification", comment="Приложите пример отчёта").json()

    mine = applicant.get(f"/api/requests/{request['id']}").json()
    assert mine["sla"]["state"] == "paused"
    assert {action["kind"] for action in mine["actions"]} == {"resubmit", "withdraw"}

    freeze(date(2026, 10, 5))
    assert admin.get("/api/requests", params={"scope": "all", "attention": "overdue"}).json()["total"] == 0

    # Черновик ответа живёт на сервере: уход со страницы ничего не теряет.
    draft = applicant.get(f"/api/drafts/{request['id']}").json()
    assert draft["version"] == 0
    saved = applicant.put(f"/api/drafts/{request['id']}", json={"data": {"process": "Новый текст"}, "version": 0}).json()
    assert saved["version"] == 1
    assert applicant.get(f"/api/drafts/{request['id']}").json()["data"] == {"process": "Новый текст"}

    body = payload(process="Экономист собирает данные из пяти отчётов; пример отчёта приложен.")
    body["version"] = mine["version"]
    response = applicant.post(f"/api/requests/{request['id']}/resubmit", json=body)
    assert response.status_code == 200, response.text
    resubmitted = response.json()

    assert resubmitted["status"] == "review"
    assert resubmitted["review_due_date"] == "2026-10-12"
    assert resubmitted["is_unviewed"] is True
    assert resubmitted["events"][-1]["kind"] == "resubmitted"
    assert resubmitted["events"][-1]["payload"]["changed"] == ["Текущий процесс"]
    assert applicant.get(f"/api/drafts/{request['id']}").status_code == 409
    assert any("дополнена" in item["text"] for item in admin.get("/api/notifications").json()["items"])


def test_withdraw_only_before_decision(applicant, admin):
    request = create_request(applicant)
    withdrawn = applicant.post(f"/api/requests/{request['id']}/withdraw", json={"comment": "Решили своими силами", "version": request["version"]})
    assert withdrawn.status_code == 200
    assert withdrawn.json()["status"] == "withdrawn"
    assert applicant.post(f"/api/requests/{request['id']}/comments", json={"text": "Ещё вопрос"}).status_code == 409

    in_work = accept(admin, create_request(applicant))
    response = applicant.post(f"/api/requests/{in_work['id']}/withdraw", json={"version": in_work["version"]})
    assert response.status_code == 409
    assert response.json()["code"] == "cannot_withdraw"


def test_internal_notes_are_hidden_from_applicant(applicant, admin):
    # АР-5: переписка по заявке без смены статуса, служебные заметки видны только администраторам.
    request = create_request(applicant)
    assert admin.post(f"/api/requests/{request['id']}/comments", json={"text": "Похоже на задачу из плана Q4", "is_internal": True}).status_code == 201
    assert admin.post(f"/api/requests/{request['id']}/comments", json={"text": "Сколько отчётов в месяц?"}).status_code == 201

    comments = [event for event in applicant.get(f"/api/requests/{request['id']}").json()["events"] if event["kind"] == "comment"]
    assert [comment["comment"] for comment in comments] == ["Сколько отчётов в месяц?"]
    assert applicant.post(f"/api/requests/{request['id']}/comments", json={"text": "x", "is_internal": True}).status_code == 403

    admin_unread_before = admin.get("/api/notifications").json()["unread"]
    assert applicant.post(f"/api/requests/{request['id']}/comments", json={"text": "Пять отчётов"}).status_code == 201
    assert admin.get("/api/notifications").json()["unread"] == admin_unread_before + 1


def test_notifications_can_be_marked_read(applicant, admin):
    create_request(applicant)
    items = admin.get("/api/notifications").json()["items"]
    assert admin.post(f"/api/notifications/{items[0]['id']}/read").status_code == 204
    assert admin.get("/api/notifications").json()["unread"] == 0
    assert applicant.post(f"/api/notifications/{items[0]['id']}/read").status_code == 204
    assert admin.post("/api/notifications/read-all").status_code == 204


# ---------------------------------------------------------------- дашборд

def test_dashboard_groups_add_up_to_total(applicant, admin, freeze):
    # Л-9: плитки дашборда складываются в общее число, отклонённые и отозванные учтены.
    stays_in_review = create_request(applicant)
    transition(admin, create_request(applicant), "rejected", comment="Нет")
    withdrawn = create_request(applicant)
    applicant.post(f"/api/requests/{withdrawn['id']}/withdraw", json={"version": withdrawn["version"]})
    accept(admin, create_request(applicant))

    dashboard = admin.get("/api/admin/dashboard").json()
    counts = {group["key"]: group["count"] for group in dashboard["groups"]}
    assert sum(counts.values()) == dashboard["total"] == 4
    assert counts["rejected"] == 1 and counts["withdrawn"] == 1 and counts["in_work"] == 1 and counts["review"] == 1
    assert dashboard["decision"]["count"] == 2
    assert dashboard["workload"]["with_estimate"] == 2

    freeze(date(2026, 9, 22))
    dashboard = admin.get("/api/admin/dashboard").json()
    assert dashboard["attention"]["overdue"] == 1
    overdue = admin.get("/api/requests", params={"scope": "all", "attention": "overdue"}).json()["items"]
    assert [item["id"] for item in overdue] == [stays_in_review["id"]]
    assert overdue[0]["sla"] == {"state": "overdue", "review_due_date": "2026-09-21", "business_days_left": -1}
    assert applicant.get("/api/admin/dashboard").status_code == 403


# ---------------------------------------------------------------- черновики и файлы

def test_drafts_are_versioned_and_discard_really_discards(applicant):
    # Д-6: «Начать заново» раньше возвращало удалённые данные.
    assert applicant.get("/api/drafts/new").status_code == 404
    first = applicant.put("/api/drafts/new", json={"data": {"process": "черновик"}}).json()
    assert first["version"] == 1

    conflict = applicant.put("/api/drafts/new", json={"data": {"process": "другая вкладка"}})
    assert conflict.status_code == 409
    assert applicant.put("/api/drafts/new", json={"data": {"process": "черновик 2"}, "version": 1}).json()["version"] == 2

    assert applicant.delete("/api/drafts/new").status_code == 204
    assert applicant.get("/api/drafts/new").status_code == 404


def test_file_upload_is_validated_attached_and_access_controlled(applicant, other_applicant, admin, monkeypatch):
    # Д-3, Д-11: файлы действительно сохраняются, тип и размер проверяются, скачать может только автор и администратор.
    from app.config import get_settings

    def upload(name, content):
        return applicant.post("/api/drafts/new/documents", files={"file": (name, content, "application/octet-stream")})

    rejected = upload("акт.<IFRAME SRC=//evil>", b"x")
    assert rejected.status_code == 422
    assert "не принимается" in rejected.json()["message"]
    assert upload("пустой.pdf", b"").status_code == 422

    monkeypatch.setattr(get_settings(), "max_upload_mb", 1)
    assert upload("большой.pdf", b"0" * (1024 * 1024 + 1)).status_code == 413

    kept = upload("../../etc/форма отчёта.pdf", b"%PDF-1.4 demo").json()
    assert kept["original_name"] == "форма отчёта.pdf"
    assert kept["extension"] == "pdf"
    extra = upload("лишний.docx", b"docx").json()
    assert len(applicant.get("/api/drafts/new").json()["documents"]) == 2

    request = create_request(applicant, documents=[{"id": kept["id"], "usage": "Столбцы 1–3", "future_use": "Заполнять автоматически"}])
    assert [(d["original_name"], d["usage"]) for d in request["documents"]] == [("форма отчёта.pdf", "Столбцы 1–3")]

    from app.db import SessionLocal
    from app.models import RequestDocument

    with SessionLocal() as db:
        assert db.get(RequestDocument, extra["id"]) is None
        stored = db.get(RequestDocument, kept["id"])
        assert storage_path(stored.storage_key).exists()

    download = applicant.get(f"/api/documents/{kept['id']}/download")
    assert download.status_code == 200
    assert download.content == b"%PDF-1.4 demo"
    assert "filename*=UTF-8''" in download.headers["content-disposition"]
    assert other_applicant.get(f"/api/documents/{kept['id']}/download").status_code == 404
    assert admin.get(f"/api/documents/{kept['id']}/download").status_code == 200


def test_attaching_foreign_document_is_rejected(applicant, other_applicant):
    foreign = other_applicant.post("/api/drafts/new/documents", files={"file": ("чужой.pdf", b"x", "application/pdf")}).json()
    response = applicant.post("/api/requests", json=payload(documents=[{"id": foreign["id"]}]))
    assert errors_by_path(response) == {"documents.0": "Файл не найден. Загрузите его заново"}


def test_resubmit_can_remove_and_add_files(applicant, admin):
    doc = applicant.post("/api/drafts/new/documents", files={"file": ("старый.xlsx", b"old", "application/octet-stream")}).json()
    request = create_request(applicant, documents=[{"id": doc["id"]}])
    request = transition(admin, request, "clarification", comment="Нужен другой пример").json()

    new_doc = applicant.post(f"/api/drafts/{request['id']}/documents", files={"file": ("новый.xlsx", b"new", "application/octet-stream")}).json()
    assert applicant.delete(f"/api/drafts/{request['id']}/documents/{doc['id']}").status_code == 404  # приложенный к заявке удаляется только при отправке

    body = payload(documents=[{"id": new_doc["id"]}])
    body["version"] = request["version"]
    result = applicant.post(f"/api/requests/{request['id']}/resubmit", json=body).json()
    assert [d["original_name"] for d in result["documents"]] == ["новый.xlsx"]
    assert result["events"][-1]["payload"]["files"] == ["добавлен файл «новый.xlsx»", "удалён файл «старый.xlsx»"]


# ---------------------------------------------------------------- безопасность и выгрузка

def test_unsafe_requests_require_csrf_header(client):
    response = client.post("/api/requests", json=payload(), headers={"X-Dev-User": "o.smirnova"})
    assert response.status_code == 403
    assert response.json()["code"] == "csrf"


def test_security_headers(applicant):
    response = applicant.get("/api/me")
    assert response.headers["x-content-type-options"] == "nosniff"
    assert "frame-ancestors" in response.headers["content-security-policy"]
    assert response.headers["cache-control"] == "no-store"


def test_excel_export_neutralises_formulas(applicant, admin):
    create_request(applicant, process='=HYPERLINK("http://evil.example","Нажмите")')
    response = admin.get("/api/requests/export.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    header = [cell.value for cell in sheet[1]]
    process = sheet.cell(row=2, column=header.index("Текущий процесс") + 1).value
    assert process.startswith("'=")
    assert applicant.get("/api/requests/export.xlsx").status_code == 403
